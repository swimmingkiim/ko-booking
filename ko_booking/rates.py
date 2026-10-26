"""Discover and validate the NTS annual expense-rate attachment."""

import re
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from io import BytesIO
from urllib.parse import urljoin
from zipfile import BadZipFile

from openpyxl import load_workbook

from .common import CollectionError, iso_date

DEFAULT_NOTICE = (
    "https://www.nts.go.kr/nts/na/ntt/selectNttInfo.do?"
    "bbsId=1143&mi=2207&nttSn=1350751"
)
HEADERS = ["귀속연도", "업종코드", "업태명", "중분류", "세분류", "세세분류",
           "적용기준내용", "단순경비율(일반율)", "단순경비율(초과율)", "기준경비율(일반율)"]
FIELDS = ["tax_year", "industry_code", "sector", "division", "class", "subclass",
          "scope", "simple_general_pct", "simple_excess_pct", "standard_general_pct"]


class NoticeParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.current = None
        self.depth = 0
        self.heading = []
        self.text = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.current = [dict(attrs).get("href", ""), ""]
        if tag == "h3":
            self.depth += 1

    def handle_data(self, data):
        self.text.append(data)
        if self.current is not None:
            self.current[1] += data
        if self.depth:
            self.heading.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.current is not None:
            self.links.append(self.current)
            self.current = None
        if tag == "h3":
            self.depth = max(0, self.depth - 1)


def discover_notice(body, url, year):
    try:
        html = body.decode("utf-8")
    except UnicodeDecodeError:
        raise CollectionError("NTS notice is not valid UTF-8") from None
    parser = NoticeParser()
    parser.feed(html)
    title = " ".join(" ".join(parser.heading).split())
    if f"{year}년 귀속" not in title or "경비율" not in title:
        raise CollectionError("Notice title does not match the requested tax year")
    attachments = [(urljoin(url, href), " ".join(name.split())) for href, name in parser.links
                   if name.strip().endswith(".xlsx") and "nttFileDownload.do?" in href
                   and f"{year}년" in name]
    if len(attachments) != 1:
        raise CollectionError("Expected exactly one matching XLSX attachment")
    published = re.search(r"작성일자\s*(\d{4}\.\d{2}\.\d{2})", " ".join(parser.text))
    if not published:
        raise CollectionError("Notice publication date is missing")
    return {"title": title, "published_at": iso_date(published[1]),
            "attachment_url": attachments[0][0], "attachment_name": attachments[0][1]}


def percent(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value))
        if not number.is_finite() or not 0 <= number <= 100:
            raise InvalidOperation
    except InvalidOperation:
        raise CollectionError("Expense rate must be a number from 0 through 100") from None
    return format(number.normalize(), "f")


def parse_rows(rows, year, sheet):
    rows = iter(rows)
    header = next(rows, ())
    if [re.sub(r"\s", "", str(c or "")) for c in header] != HEADERS:
        raise CollectionError("Unexpected expense-rate columns; review the new source schema")
    records, seen = [], set()
    for row_number, row in enumerate(rows, 2):
        if all(c is None for c in row):
            continue
        if len(row) != len(FIELDS):
            raise CollectionError(f"Unexpected column count at row {row_number}")
        if str(row[0]) != str(year):
            raise CollectionError(f"Tax year mismatch at row {row_number}")
        # Never turn a numeric code into a guessed, zero-padded identifier.
        code = row[1]
        if not isinstance(code, str) or not re.fullmatch(r"\d{6}", code):
            raise CollectionError(f"Expected a six-character text industry code at row {row_number}")
        if code in seen:
            raise CollectionError(f"Duplicate industry code at row {row_number}")
        seen.add(code)
        record = dict(zip(FIELDS, row))
        record["tax_year"] = year
        for field in FIELDS[7:]:
            record[field] = percent(record[field])
        record["source_sheet"] = sheet
        record["source_row"] = row_number
        records.append(record)
    if not records:
        raise CollectionError("Expense-rate workbook contains no records")
    return records


def parse_workbook(body, year):
    try:
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=False)
        try:
            if len(workbook.worksheets) != 1:
                raise CollectionError("Expected one expense-rate sheet; review changed layout")
            sheet = workbook.worksheets[0]
            return parse_rows(sheet.iter_rows(values_only=True), year, sheet.title)
        finally:
            workbook.close()
    except (BadZipFile, KeyError):
        raise CollectionError("Attachment is not a valid XLSX workbook") from None


def collect_rates(client, archive, year, notice_url=DEFAULT_NOTICE):
    page, page_meta = client.get(notice_url)
    notice = discover_notice(page, notice_url, year)
    body, meta = client.get(notice["attachment_url"])
    records = parse_workbook(body, year)
    normalized = {"schema_version": 1, "tax_year": year, "rate_unit": "percent_points",
                  "records": records}
    # Preserve the discovery notice too, but never publish it as a complete rate dataset.
    archive.save(f"nts:expense-rates:{year}:notice", page,
                 {"schema_version": 1, **notice}, {**page_meta, "extension": ".html"})
    metadata = {**meta, **notice, "notice_url": notice_url, "tax_year": year,
                "effective_at": None, "publisher": "국세청", "extension": ".xlsx",
                "record_count": len(records), "review_status": "source_parsed_not_tax_rule_validated",
                "reuse_status": "not_confirmed", "redistribution": "not_in_repository",
                "terms_url": "https://www.nts.go.kr/nts/cm/cntnts/cntntsView.do?mi=6791&cntntsId=8169"}
    return {**archive.save(f"nts:expense-rates:{year}", body, normalized, metadata),
            "record_count": len(records)}
